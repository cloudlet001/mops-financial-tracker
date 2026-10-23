"""
金融股月自結損益彙整
1. 抓最新月份 月自結損益清單 (上市)，擷取所有金融股 → 「自結清單」
2. 對指定個股抓 個股公告自結損益表 (上一年度 + 今年度)，依月份對齊 → 每檔一個分頁
3. 將最新月份資料 (前一年本月/累計稅前、今年本月/累計稅前、累計稅前YoY) 回填至自結清單
"""
import os
import re
import sys
import json
import time
import datetime
import urllib.request

import pandas as pd
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
TARGET_CODES = ["2881", "2882", "2891", "2884", "2885"]
MARKET = "1"  # 上市
FIN_MARK = "不適用"
API_URL = "https://mops.twse.com.tw/mops/api/redirectToOld"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36",
    "Content-Type": "application/json",
    "Referer": "https://mops.twse.com.tw/mops/",
    "Origin": "https://mops.twse.com.tw",
}
DELAY = 1.5


# ---------------------------------------------------------------- helpers
def mops_get(api_name, params, retries=3):
    body = json.dumps({"apiName": api_name, "parameters": params}).encode("utf-8")
    last_err = None
    for i in range(retries):
        try:
            req = urllib.request.Request(API_URL, data=body, headers=HEADERS, method="POST")
            with urllib.request.urlopen(req, timeout=30) as r:
                res = json.loads(r.read().decode("utf-8"))
            if res.get("code") != 200:
                raise RuntimeError(f"API error: {res}")
            url = res["result"]["url"]
            hreq = urllib.request.Request(url, headers={"User-Agent": HEADERS["User-Agent"], "Referer": HEADERS["Referer"]})
            with urllib.request.urlopen(hreq, timeout=30) as r:
                html = r.read().decode("utf-8", errors="ignore")
            time.sleep(DELAY)
            return html
        except Exception as e:
            last_err = e
            time.sleep(DELAY * (i + 2))
    raise RuntimeError(f"{api_name} failed: {last_err}")


def to_num(s):
    if s is None:
        return None
    s = str(s).strip().replace(",", "")
    if s == "" or FIN_MARK in s:
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return float(s)
        except ValueError:
            return None


def norm_label(s):
    s = s.strip()
    for t in ["（仟元）：", "（最多60個中文字）：", "：", "合併"]:
        s = s.replace(t, "")
    return s


def yoy(cur, prev):
    if cur is None or prev is None or pd.isna(cur) or pd.isna(prev) or prev == 0:
        return None
    return (cur - prev) / abs(prev)


# ---------------------------------------------------------------- step 1
def fetch_list(ym):
    html = mops_get("ajax_t138sb01", {
        "CK2": "1", "YM1": ym, "CK1": MARKET,
        "encodeURIComponent": 1, "firstin": 1, "step": 1, "first": 1,
    })
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    for tbl in soup.find_all("table"):
        ths = tbl.find_all("th")
        if not ths or "公司代號" not in ths[0].text:
            continue
        # report type from the preceding caption table
        prev = tbl.find_previous("table", class_="noBorder")
        rtype = "合併" if prev and "合併" in prev.text else "個別"
        headers = []
        for th in ths:
            h = norm_label(th.text)
            if h not in headers:
                headers.append(h)
        for tr in tbl.find_all("tr"):
            tds = tr.find_all("td")
            if len(tds) != len(headers):
                continue
            rec = {h: td.text.strip().replace("\xa0", " ") for h, td in zip(headers, tds)}
            rec["報表類型"] = rtype
            rows.append(rec)
    return pd.DataFrame(rows)


def find_latest_ym():
    today = datetime.date.today()
    y, m = today.year - 1911, today.month
    for _ in range(6):
        m -= 1
        if m == 0:
            y, m = y - 1, 12
        ym = f"{y}{m:02d}"
        print(f"[*] 嘗試 {ym} ...")
        df = fetch_list(ym)
        if not df.empty:
            return ym, df
    raise RuntimeError("找不到最近 6 個月的月自結資料")


# ---------------------------------------------------------------- step 2
def fetch_company_year(co_id, year):
    q1 = mops_get("ajax_t138sb02_q1", {
        "co_id": co_id, "year": str(year),
        "encodeURIComponent": 1, "step": 1, "firstin": 1, "off": 1,
    })
    hidden = dict(re.findall(r"name='(\w+)' value='([^']*)'", q1))
    params = {"encodeURIComponent": 1, "step": 1, "firstin": 1, "off": 1}
    for k in ["CK2", "BK1", "TYPEK", "YEAR", "COMP"]:
        if k in hidden:
            params[k] = hidden[k]
    html = mops_get("ajax_t138sb02", params)
    soup = BeautifulSoup(html, "html.parser")

    name = ""
    h3s = soup.find_all("h3")
    for h in h3s:
        mm = re.search(r"公司名稱：\s*(\S+)", h.text)
        if mm:
            name = mm.group(1)

    caps = [t for t in soup.find_all("table", class_="noBorder") if re.search(rf"{year}年\s*\d+月", t.text)]
    datas = soup.find_all("table", class_="hasBorder")
    recs = []
    for cap, data in zip(caps, datas):
        month = int(re.search(rf"{year}年\s*(\d+)月", cap.text).group(1))
        rec = {"月份_num": month}
        for tr in data.find_all("tr"):
            th, td = tr.find("th"), tr.find("td")
            if th and td:
                rec[norm_label(th.text)] = td.text.strip()
        recs.append(rec)
    return name, pd.DataFrame(recs)


def build_company_sheet(co_id, prev_y, cur_y):
    name, d_prev = fetch_company_year(co_id, prev_y)
    name2, d_cur = fetch_company_year(co_id, cur_y)
    name = name or name2
    base = pd.DataFrame({"月份_num": range(1, 13)})
    for y, d in [(prev_y, d_prev), (cur_y, d_cur)]:
        sub = pd.DataFrame({"月份_num": d["月份_num"]}) if not d.empty else pd.DataFrame({"月份_num": []})
        for k in ["本月稅前損益", "本年累計稅前損益"]:
            sub[f"{y}年度 {k}"] = d[k].map(to_num) if k in d else None
        base = base.merge(sub, on="月份_num", how="left")
    base["單月稅前 YoY"] = [yoy(c, p) for c, p in zip(base[f"{cur_y}年度 本月稅前損益"], base[f"{prev_y}年度 本月稅前損益"])]
    base["累計稅前 YoY"] = [yoy(c, p) for c, p in zip(base[f"{cur_y}年度 本年累計稅前損益"], base[f"{prev_y}年度 本年累計稅前損益"])]
    base.insert(0, "月份", base["月份_num"].map(lambda m: f"{m}月"))
    return name, base


# ---------------------------------------------------------------- excel styling
def style_sheet(ws, pct_cols=(), num_start_col=2):
    hdr_fill = PatternFill("solid", start_color="1F497D")
    hdr_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
    font = Font(name="微軟正黑體", size=10)
    border = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
    headers = [c.value for c in ws[1]]
    for c in ws[1]:
        c.fill, c.font = hdr_fill, hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = font, border
            h = headers[c.column - 1]
            if h in pct_cols:
                c.number_format = '0.00%;[Red]-0.00%'
                c.alignment = Alignment(horizontal="right")
            elif isinstance(c.value, (int, float)):
                c.number_format = '#,##0;[Red]-#,##0'
                c.alignment = Alignment(horizontal="right")
            else:
                c.alignment = Alignment(horizontal="center" if c.column == 1 else "left")
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(max(width * 1.6, 12), 40)
    ws.freeze_panes = "B2"


# ---------------------------------------------------------------- main
def main():
    ym, df_list = find_latest_ym()
    cur_y, cur_m = int(ym[:3]), int(ym[3:])
    prev_y = cur_y - 1
    print(f"[+] 最新月份: {cur_y}年{cur_m}月，清單共 {len(df_list)} 筆")

    fin = df_list[df_list["本月營業損益"].str.contains(FIN_MARK, na=False)].copy()
    fin = fin.sort_values("公司代號").reset_index(drop=True)
    print(f"[+] 金融股共 {len(fin)} 檔")

    summary = pd.DataFrame({
        "金融股": fin["公司代號"] + " " + fin["公司簡稱"],
        "報表類型": fin["報表類型"],
        f"{cur_y}/{cur_m:02d} 本月稅前損益(清單)": fin["本月稅前損益"].map(to_num),
        f"{cur_y}/{cur_m:02d} 本年累計稅前損益(清單)": fin["本年累計稅前損益"].map(to_num),
    })
    fill_cols = [f"{prev_y}年度 本月稅前損益", f"{prev_y}年度 本年累計稅前損益",
                 f"{cur_y}年度 本月稅前損益", f"{cur_y}年度 本年累計稅前損益", "累計稅前 YoY"]
    for c in fill_cols:
        summary[c] = None

    company_sheets = {}
    for co in TARGET_CODES:
        print(f"[*] 抓取 {co} 的 {prev_y}/{cur_y} 個股自結損益 ...")
        name, sheet = build_company_sheet(co, prev_y, cur_y)
        company_sheets[f"{co} {name}"] = sheet
        latest = sheet[sheet["月份_num"] == cur_m].iloc[0]
        idx = fin.index[fin["公司代號"] == co]
        if len(idx) == 0:
            print(f"    [!] {co} 不在 {ym} 金融股清單中，新增一列")
            summary.loc[len(summary)] = {"金融股": f"{co} {name}", "報表類型": ""}
            idx = [len(summary) - 1]
        i = idx[0]
        for c in fill_cols:
            summary.at[i, c] = latest[c]
        # cross-check
        lst_m = summary.at[i, f"{cur_y}/{cur_m:02d} 本月稅前損益(清單)"]
        lst_c = summary.at[i, f"{cur_y}/{cur_m:02d} 本年累計稅前損益(清單)"]
        ok = (lst_m == latest[f"{cur_y}年度 本月稅前損益"]) and (lst_c == latest[f"{cur_y}年度 本年累計稅前損益"])
        print(f"    {co} {name}: 交叉驗證 {'OK' if ok else 'MISMATCH'}  "
              f"(清單 {lst_m}/{lst_c} vs 個股 {latest[f'{cur_y}年度 本月稅前損益']}/{latest[f'{cur_y}年度 本年累計稅前損益']})")

    # write
    xlsx = os.path.join(OUTPUT_DIR, f"financial_self_profit_{ym}.xlsx")
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        summary.to_excel(w, sheet_name="自結清單", index=False)
        for sname, sheet in company_sheets.items():
            sheet.drop(columns=["月份_num"]).to_excel(w, sheet_name=sname[:31], index=False)
    wb = openpyxl.load_workbook(xlsx)
    for ws in wb.worksheets:
        style_sheet(ws, pct_cols=("累計稅前 YoY", "單月稅前 YoY"))
    wb.save(xlsx)
    summary.to_csv(os.path.join(OUTPUT_DIR, f"financial_self_profit_{ym}_list.csv"), index=False, encoding="utf-8-sig")

    pd.set_option("display.width", 250)
    print("\n=== 自結清單 ===")
    print(summary.to_string())
    for sname, sheet in company_sheets.items():
        print(f"\n=== {sname} ===")
        print(sheet.drop(columns=["月份_num"]).to_string())
    print(f"\n[+] 輸出: {xlsx}")


if __name__ == "__main__":
    main()
