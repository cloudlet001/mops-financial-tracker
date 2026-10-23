import sys
import os
import pandas as pd
from bs4 import BeautifulSoup
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_DIR = r"C:\Users\USER\.gemini\antigravity\scratch\mops_monthly_profit"

def parse_year(year):
    html_file = os.path.join(OUTPUT_DIR, f"2881_{year}_full.html")
    with open(html_file, "r", encoding="utf-8") as f:
        html = f.read()
    
    soup = BeautifulSoup(html, 'html.parser')
    no_border_tables = [t for t in soup.find_all('table', class_='noBorder') if f'{year}年' in t.text]
    has_border_tables = soup.find_all('table', class_='hasBorder')
    
    records = []
    for nb, hb in zip(no_border_tables, has_border_tables):
        m_match = re.search(rf'{year}年\s*(\d+)月', nb.text)
        month_int = int(m_match.group(1)) if m_match else None
        
        row_dict = {
            "月份_num": month_int,
            "月份": f"{month_int}月"
        }
        for tr in hb.find_all('tr'):
            th = tr.find('th')
            td = tr.find('td')
            if th and td:
                key = th.text.strip().replace('（仟元）：', '').replace('（最多60個中文字）：', '').replace('：', '')
                val = td.text.strip().replace('\xa0', ' ')
                row_dict[key] = val
        records.append(row_dict)
    
    df = pd.DataFrame(records)
    return df

df_114 = parse_year("114")
df_115 = parse_year("115")

# Create full 1-12 months base
months_base = pd.DataFrame({
    "月份_num": range(1, 13),
    "月份": [f"{m}月" for m in range(1, 13)]
})

# Prefix columns with year
cols_to_keep = ['本月合併營業損益', '本年累計合併營業損益', '本月合併稅前損益', '本年累計合併稅前損益']

df_114_sub = df_114[['月份_num'] + cols_to_keep].copy()
df_114_sub.columns = ['月份_num'] + [f"114年度_{c}" for c in cols_to_keep]

df_115_sub = df_115[['月份_num'] + cols_to_keep].copy()
df_115_sub.columns = ['月份_num'] + [f"115年度_{c}" for c in cols_to_keep]

# Merge by month
merged = months_base.merge(df_114_sub, on='月份_num', how='left').merge(df_115_sub, on='月份_num', how='left')
merged = merged.drop(columns=['月份_num'])
merged = merged.fillna('')

# Also create a key metrics comparison view
# Monthly pre-tax profit & Cumulative pre-tax profit
comp_df = months_base[['月份_num', '月份']].copy()
# convert numbers
for y, d in [('114', df_114), ('115', df_115)]:
    d_clean = d[['月份_num', '本月合併稅前損益', '本年累計合併稅前損益']].copy()
    d_clean[f'{y}年_單月稅前損益(仟元)'] = d_clean['本月合併稅前損益'].str.replace(',', '').astype(float)
    d_clean[f'{y}年_累計稅前損益(仟元)'] = d_clean['本年累計合併稅前損益'].str.replace(',', '').astype(float)
    d_clean = d_clean.drop(columns=['本月合併稅前損益', '本年累計合併稅前損益'])
    comp_df = comp_df.merge(d_clean, on='月份_num', how='left')

# Calculate YoY for single month pre-tax
comp_df['單月YoY成長率'] = ((comp_df['115年_單月稅前損益(仟元)'] - comp_df['114年_單月稅前損益(仟元)']) / comp_df['114年_單月稅前損益(仟元)'].abs() * 100).round(2)
comp_df['累計YoY成長率'] = ((comp_df['115年_累計稅前損益(仟元)'] - comp_df['114年_累計稅前損益(仟元)']) / comp_df['114年_累計稅前損益(仟元)'].abs() * 100).round(2)

comp_df_display = comp_df.drop(columns=['月份_num'])

print("--- 完整跨年度月份對齊表 ---")
print(merged.to_string())

# Save to CSV
csv_path = os.path.join(OUTPUT_DIR, "fubon_2881_monthly_profit_114_115.csv")
merged.to_csv(csv_path, index=False, encoding='utf-8-sig')

# Save to Excel
excel_path = os.path.join(OUTPUT_DIR, "fubon_2881_monthly_profit_114_115.xlsx")
with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    merged.to_excel(writer, sheet_name='自結損益彙整表(114-115)', index=False)
    comp_df_display.to_excel(writer, sheet_name='稅前損益YoY比較', index=False)

# Style Excel
wb = openpyxl.load_workbook(excel_path)
for ws in wb.worksheets:
    ws.views.sheetView[0].showGridLines = True
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="微軟正黑體", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            c = ws.cell(row=row, column=col)
            c.font = cell_font
            c.border = thin_border
            if col == 1:
                c.alignment = Alignment(horizontal="center")
            else:
                c.alignment = Alignment(horizontal="right")
    
    # Auto fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len * 1.8, 14)

wb.save(excel_path)
print(f"\n[+] 儲存完成！")
print(f"CSV: {csv_path}")
print(f"Excel: {excel_path}")
