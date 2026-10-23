"""Build the Excel deliverables:
 - modMopsSelfProfit.bas : VBA module with all non-ASCII string literals converted to U("\\uXXXX")
                           so it imports correctly regardless of Windows locale / code page.
 - MOPS_金融股自結損益.xlsx : template workbook with the 設定 sheet (save as .xlsm after importing the module).
"""
import os
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "modMopsSelfProfit.src.bas")
OUT_BAS = os.path.join(HERE, "modMopsSelfProfit.bas")
OUT_XLSX = os.path.join(HERE, "MOPS_金融股自結損益.xlsx")

LIT = re.compile(r'"((?:[^"\r\n]|"")*)"')


def conv_literal(m):
    body = m.group(1)
    if all(ord(ch) < 128 for ch in body):
        return m.group(0)
    esc = "".join(ch if ord(ch) < 128 else "\\u%04x" % ord(ch) for ch in body)
    return 'U("%s")' % esc


def build_bas():
    src = open(SRC, encoding="utf-8").read()
    out_lines = []
    for line in src.splitlines():
        stripped = line.lstrip()
        if stripped.startswith("'"):
            out_lines.append(line)
            continue
        out_lines.append(LIT.sub(conv_literal, line))
    out = "\r\n".join(out_lines) + "\r\n"
    bad = [(i + 1, l) for i, l in enumerate(out.splitlines()) if any(ord(c) > 127 for c in l)]
    if bad:
        raise SystemExit(f"non-ASCII left in output: {bad[:5]}")
    long_lines = [i + 1 for i, l in enumerate(out.splitlines()) if len(l) > 1000]
    if long_lines:
        raise SystemExit(f"lines too long for VBA: {long_lines}")
    open(OUT_BAS, "w", encoding="ascii", newline="").write(out)
    print("wrote", OUT_BAS)


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "設定"
    f = "微軟正黑體"
    ws["A1"] = "MOPS 金融股月自結損益彙整"
    ws["A1"].font = Font(name=f, size=16, bold=True, color="1F497D")
    labels = {
        "A3": "市場別", "C3": "1=上市  2=上櫃  3=興櫃",
        "A4": "年月 (民國 YYYMM)", "C4": "空白 = 自動抓最新月份，例: 11509",
        "A6": "目標個股代號", "C6": "由 B7 往下填，遇空白停止",
        "D3": "執行狀態", "D5": "執行紀錄",
        "A20": "執行方式: 按 Alt+F8 → 選 RunFinancialSelfProfit → 執行 (或按設定好的按鈕)",
        "A21": "每次執行會刪除並重建「設定」以外的所有分頁。",
    }
    for k, v in labels.items():
        ws[k] = v
        ws[k].font = Font(name=f, size=10, bold=k in ("A3", "A4", "A6", "D3", "D5"),
                          color="808080" if k.startswith("C") or k in ("A20", "A21") else "000000")
    ws["B3"] = "1"
    ws["B4"] = ""
    for c in ("B3", "B4"):
        ws[c].number_format = "@"
    for i, code in enumerate(["2881", "2882", "2891", "2884", "2885"]):
        cell = ws.cell(row=7 + i, column=2, value=code)
        cell.number_format = "@"
    inp = PatternFill("solid", start_color="FFF2CC")
    thin = Side(style="thin", color="BFBFBF")
    for r in list(range(3, 5)) + list(range(7, 18)):
        c = ws.cell(row=r, column=2)
        c.fill = inp
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        c.font = Font(name=f, size=11)
        c.alignment = Alignment(horizontal="center")
        c.number_format = "@"
    ws["E3"].font = Font(name=f, size=10, bold=True, color="C00000")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 50
    wb.save(OUT_XLSX)
    print("wrote", OUT_XLSX)


if __name__ == "__main__":
    build_bas()
    build_xlsx()
