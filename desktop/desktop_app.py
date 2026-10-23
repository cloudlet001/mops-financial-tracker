"""
MOPS 金融股月自結損益彙整系統 — Windows 桌面應用程式版 (Native Tkinter/TTK GUI)
完整 Clone 自 Web App，具備：
1. 開啟時秒速載入最新「自結清單」(優先從 SQLite3 載入，不自動抓個股，由使用者勾選後更新或一鍵更新 ALL)
2. 頂部標題列整合「📋 自結清單」與動態個股分頁 (含單一「×」關閉與「🗑️ 清除所有個股分頁」)
3. 左上角卡片內建 3 個分頁 (Tabs)：
   - ① 清單年月 & 兩年度選擇 (預設上一年度 vs 今年度)
   - 📋 自結清單 資料來源 (顯示 POST URL & Payload，滑鼠點擊直接開啟 MOPS 原始報表)
   - 📈 個股自結損益表 資料來源 (顯示 POST URL & Payload，支援切換個股並點擊直接開啟 MOPS 報表)
4. 右上角卡片提供金融股勾選清單、預設五檔/前三檔/全選/清除、以及「🔄 強制重抓」(左) vs 「🗄️ 資料庫優先」(右) 切換
5. SQLite3 資料庫持久化 (mops_self_profit.db，與 Web 版完全相容共用 Schema)
6. 自結清單全部 10 個欄位皆支援點擊表頭排序 (▲ / ▼，無資料列自動置底)，表頭在「本」字換行避免水平過寬
7. 支援匯出美化格式 Excel (.xlsx) 與清單 CSV
"""
import os
import re
import sys
import json
import time
import shutil
import sqlite3
import datetime
import threading
import webbrowser
import urllib.parse
import urllib.request
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

import pandas as pd
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(HERE, "mops_self_profit.db")
WEB_DB_PATH = os.path.join(os.path.dirname(HERE), "web", "mops_self_profit.db")
DB_LOCK = threading.Lock()

FIN_MARK = "不適用"
API_URL = "https://mops.twse.com.tw/mops/api/redirectToOld"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36",
    "Content-Type": "application/json",
    "Referer": "https://mops.twse.com.tw/mops/",
    "Origin": "https://mops.twse.com.tw",
}
DELAY = 0.8


# ================================================================ SQLite3 資料庫層 (與 web_app.py 完全相同 Schema)
def init_db():
    # 若桌面版 DB 不存在或無 mops_list 資料表，優先從 web 目錄同步最新 mops_self_profit.db
    need_copy = not os.path.exists(DB_PATH)
    if not need_copy and os.path.exists(WEB_DB_PATH):
        try:
            with sqlite3.connect(DB_PATH) as c:
                tbl = c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='mops_list'").fetchone()
                if not tbl:
                    need_copy = True
                else:
                    cnt = c.execute("SELECT count(*) FROM mops_list").fetchone()[0]
                    if cnt == 0:
                        need_copy = True
        except Exception:
            need_copy = True

    if need_copy and os.path.exists(WEB_DB_PATH):
        try:
            shutil.copy2(WEB_DB_PATH, DB_PATH)
        except Exception:
            pass

    with DB_LOCK, sqlite3.connect(DB_PATH) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS mops_list_meta (
                ym TEXT NOT NULL,
                market TEXT NOT NULL,
                total_companies INTEGER NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (ym, market)
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS mops_list (
                ym TEXT NOT NULL,
                market TEXT NOT NULL,
                co_id TEXT NOT NULL,
                name TEXT NOT NULL,
                report_type TEXT NOT NULL,
                list_month_pretax REAL,
                list_cum_pretax REAL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (ym, market, co_id)
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS company_year_meta (
                co_id TEXT NOT NULL,
                year INTEGER NOT NULL,
                name TEXT NOT NULL,
                months_count INTEGER NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (co_id, year)
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS company_monthly_profit (
                co_id TEXT NOT NULL,
                year INTEGER NOT NULL,
                month INTEGER NOT NULL,
                month_pretax REAL,
                cum_pretax REAL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (co_id, year, month)
            )
        """)
        conn.commit()


def now_str():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def db_get_list(ym, market):
    with DB_LOCK, sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        meta = conn.execute(
            "SELECT total_companies, updated_at FROM mops_list_meta WHERE ym = ? AND market = ?",
            (str(ym), str(market)),
        ).fetchone()
        if not meta:
            return None
        rows = conn.execute(
            """SELECT co_id, name, report_type, list_month_pretax, list_cum_pretax
               FROM mops_list WHERE ym = ? AND market = ? ORDER BY co_id ASC""",
            (str(ym), str(market)),
        ).fetchall()
        if not rows:
            return None
        stocks = []
        for r in rows:
            m_val = int(r["list_month_pretax"]) if r["list_month_pretax"] is not None else None
            c_val = int(r["list_cum_pretax"]) if r["list_cum_pretax"] is not None else None
            stocks.append({
                "co_id": r["co_id"],
                "name": r["name"],
                "label": f"{r['co_id']} {r['name']}",
                "report_type": r["report_type"],
                "list_month_pretax": m_val,
                "list_cum_pretax": c_val,
            })
        return {
            "ym": str(ym),
            "year": int(str(ym)[:-2]),
            "month": int(str(ym)[-2:]),
            "market": str(market),
            "total_companies": meta["total_companies"],
            "financial_count": len(stocks),
            "stocks": stocks,
            "source": "db",
            "updated_at": meta["updated_at"],
        }


def db_save_list(ym, market, total_companies, stocks):
    ts = now_str()
    ym, market = str(ym), str(market)
    with DB_LOCK, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """INSERT OR REPLACE INTO mops_list_meta (ym, market, total_companies, updated_at)
               VALUES (?, ?, ?, ?)""",
            (ym, market, total_companies, ts),
        )
        conn.execute("DELETE FROM mops_list WHERE ym = ? AND market = ?", (ym, market))
        conn.executemany(
            """INSERT INTO mops_list (ym, market, co_id, name, report_type, list_month_pretax, list_cum_pretax, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (
                    ym,
                    market,
                    s["co_id"],
                    s["name"],
                    s["report_type"],
                    s["list_month_pretax"],
                    s["list_cum_pretax"],
                    ts,
                )
                for s in stocks
            ],
        )
        conn.commit()
    return ts


def db_get_company_year(co_id, year):
    co_id, year = str(co_id), int(year)
    with DB_LOCK, sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        meta = conn.execute(
            "SELECT name, months_count, updated_at FROM company_year_meta WHERE co_id = ? AND year = ?",
            (co_id, year),
        ).fetchone()
        if not meta or meta["months_count"] <= 0:
            return None
        rows = conn.execute(
            """SELECT month, month_pretax, cum_pretax
               FROM company_monthly_profit WHERE co_id = ? AND year = ? ORDER BY month ASC""",
            (co_id, year),
        ).fetchall()
        if not rows:
            return None
        month_map = {}
        for r in rows:
            month_map[int(r["month"])] = {
                "month_pretax": int(r["month_pretax"]) if r["month_pretax"] is not None else None,
                "cum_pretax": int(r["cum_pretax"]) if r["cum_pretax"] is not None else None,
            }
        return {
            "name": meta["name"],
            "month_map": month_map,
            "updated_at": meta["updated_at"],
            "source": "db",
        }


def db_save_company_year(co_id, year, name, month_map):
    co_id, year = str(co_id), int(year)
    ts = now_str()
    with DB_LOCK, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """INSERT OR REPLACE INTO company_year_meta (co_id, year, name, months_count, updated_at)
               VALUES (?, ?, ?, ?, ?)""",
            (co_id, year, name or "", len(month_map), ts),
        )
        conn.execute("DELETE FROM company_monthly_profit WHERE co_id = ? AND year = ?", (co_id, year))
        conn.executemany(
            """INSERT INTO company_monthly_profit (co_id, year, month, month_pretax, cum_pretax, updated_at)
               VALUES (?, ?, ?, ?, ?, ?)""",
            [
                (co_id, year, int(m), d.get("month_pretax"), d.get("cum_pretax"), ts)
                for m, d in month_map.items()
            ],
        )
        conn.commit()
    return ts


# ================================================================ MOPS 網路爬蟲層
def resolve_mops_url(api_name: str, parameters: dict) -> str:
    body = json.dumps({"apiName": api_name, "parameters": parameters}).encode("utf-8")
    req = urllib.request.Request(API_URL, data=body, headers=HEADERS, method="POST")
    with urllib.request.urlopen(req, timeout=25) as r:
        res = json.loads(r.read().decode("utf-8"))
    if res.get("code") != 200 or "result" not in res or "url" not in res["result"]:
        raise RuntimeError(f"MOPS redirectToOld 失敗: {res}")
    return res["result"]["url"]


def mops_get(api_name, params, retries=3):
    body = json.dumps({"apiName": api_name, "parameters": params}).encode("utf-8")
    last_err = None
    for i in range(retries):
        try:
            req = urllib.request.Request(API_URL, data=body, headers=HEADERS, method="POST")
            with urllib.request.urlopen(req, timeout=25) as r:
                res = json.loads(r.read().decode("utf-8"))
            if res.get("code") != 200 or "result" not in res or "url" not in res["result"]:
                raise RuntimeError(f"MOPS API 回應非 200 或無資料: {res.get('message', res)}")
            url = res["result"]["url"]
            hreq = urllib.request.Request(
                url,
                headers={"User-Agent": HEADERS["User-Agent"], "Referer": HEADERS["Referer"]},
            )
            with urllib.request.urlopen(hreq, timeout=25) as r:
                html = r.read().decode("utf-8", errors="ignore")
            if "FOR SECURITY REASONS" in html:
                raise RuntimeError("MOPS 安全機制阻擋，稍後重試")
            time.sleep(DELAY)
            return html
        except Exception as e:
            last_err = e
            time.sleep(DELAY * (i + 1.5))
    raise RuntimeError(f"{api_name} 請求失敗: {last_err}")


def to_num(s):
    if s is None:
        return None
    s = str(s).strip().replace(",", "")
    if s == "" or FIN_MARK in s:
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return float(s)
        except ValueError:
            return None


def norm_label(s):
    s = s.strip()
    for t in ["（仟元）：", "（最多60個中文字）：", "：", "合併"]:
        s = s.replace(t, "")
    return s


def calc_yoy(cur, prev):
    if cur is None or prev is None or prev == 0:
        return None
    return (cur - prev) / abs(prev)


def fetch_list_data(ym, market="1", force=False):
    ym, market = str(ym), str(market)
    if not force:
        cached = db_get_list(ym, market)
        if cached:
            return cached

    html = mops_get("ajax_t138sb01", {
        "CK2": "1",
        "YM1": ym,
        "CK1": market,
        "encodeURIComponent": 1,
        "firstin": 1,
        "step": 1,
        "first": 1,
    })
    soup = BeautifulSoup(html, "html.parser")
    all_rows = []
    for tbl in soup.find_all("table"):
        ths = tbl.find_all("th")
        if not ths or "公司代號" not in ths[0].text:
            continue
        prev = tbl.find_previous("table", class_="noBorder")
        rtype = "合併" if prev and "合併" in prev.text else "個別"
        headers = []
        for th in ths:
            h = norm_label(th.text)
            if h not in headers:
                headers.append(h)
        for tr in tbl.find_all("tr"):
            tds = tr.find_all("td")
            if len(tds) != len(headers):
                continue
            rec = {h: td.text.strip().replace("\xa0", " ") for h, td in zip(headers, tds)}
            rec["報表類型"] = rtype
            all_rows.append(rec)

    fin_stocks = []
    for r in all_rows:
        op_val = r.get("本月營業損益", "")
        if FIN_MARK in op_val:
            fin_stocks.append({
                "co_id": r.get("公司代號", ""),
                "name": r.get("公司簡稱", ""),
                "label": f"{r.get('公司代號', '')} {r.get('公司簡稱', '')}",
                "report_type": r.get("報表類型", ""),
                "list_month_pretax": to_num(r.get("本月稅前損益")),
                "list_cum_pretax": to_num(r.get("本年累計稅前損益")),
            })

    if not fin_stocks:
        raise RuntimeError(f"年月 {ym} 查無月自結損益彙總表資料")

    fin_stocks.sort(key=lambda x: x["co_id"])
    ts = db_save_list(ym, market, len(all_rows), fin_stocks)
    return {
        "ym": ym,
        "year": int(ym[:-2]),
        "month": int(ym[-2:]),
        "market": market,
        "total_companies": len(all_rows),
        "financial_count": len(fin_stocks),
        "stocks": fin_stocks,
        "source": "mops",
        "updated_at": ts,
    }


def find_latest_ym(market="1"):
    # 1. 優先檢查 SQLite 資料庫中是否已有最新月份清單（秒開不卡頓）
    with DB_LOCK, sqlite3.connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT ym FROM mops_list_meta WHERE market = ? ORDER BY ym DESC LIMIT 1",
            (str(market),),
        ).fetchone()
    if row:
        cached = db_get_list(row[0], market)
        if cached and cached.get("stocks"):
            return cached

    # 2. 若資料庫尚無任何清單，才往回探測最近 6 個月
    today = datetime.date.today()
    y, m = today.year - 1911, today.month
    for _ in range(6):
        m -= 1
        if m == 0:
            y, m = y - 1, 12
        ym = f"{y}{m:02d}"
        try:
            res = fetch_list_data(ym, market, force=False)
            if res["total_companies"] > 0:
                return res
        except Exception:
            pass
    raise RuntimeError("找不到最近 6 個月的月自結損益清單")


def fetch_company_year_data(co_id, year, force=False):
    co_id, year = str(co_id), int(year)
    if not force:
        db_res = db_get_company_year(co_id, year)
        if db_res:
            return db_res["name"], db_res["month_map"], "db", db_res["updated_at"]

    try:
        q1 = mops_get("ajax_t138sb02_q1", {
            "co_id": co_id,
            "year": str(year),
            "encodeURIComponent": 1,
            "step": 1,
            "firstin": 1,
            "off": 1,
        })
        hidden = dict(re.findall(r"name='(\w+)' value='([^']*)'", q1))
        params = {"encodeURIComponent": 1, "step": 1, "firstin": 1, "off": 1}
        for k in ["CK2", "BK1", "TYPEK", "YEAR", "COMP"]:
            if k in hidden:
                params[k] = hidden[k]
        if "COMP" not in params:
            params["COMP"] = co_id
            params["YEAR"] = str(year)
            params["CK2"] = "1"
            params["BK1"] = "2"
            params["TYPEK"] = "sii"

        html = mops_get("ajax_t138sb02", params)
    except Exception:
        return "", {}, "mops", now_str()

    soup = BeautifulSoup(html, "html.parser")
    name = ""
    for h in soup.find_all("h3"):
        mm = re.search(r"公司名稱：\s*(\S+)", h.text)
        if mm:
            name = mm.group(1).strip()

    caps = [t for t in soup.find_all("table", class_="noBorder") if re.search(rf"{year}年\s*\d+月", t.text)]
    datas = soup.find_all("table", class_="hasBorder")
    month_map = {}
    for cap, data in zip(caps, datas):
        m_match = re.search(rf"{year}年\s*(\d+)月", cap.text)
        if not m_match:
            continue
        month = int(m_match.group(1))
        rec = {}
        for tr in data.find_all("tr"):
            th, td = tr.find("th"), tr.find("td")
            if th and td:
                rec[norm_label(th.text)] = td.text.strip()
        month_map[month] = {
            "month_pretax": to_num(rec.get("本月稅前損益")),
            "cum_pretax": to_num(rec.get("本年累計稅前損益")),
        }

    ts = db_save_company_year(co_id, year, name, month_map) if month_map else now_str()
    return name, month_map, "mops", ts


def build_company_comparison(co_id, y1, y2, ref_month=None, force=False):
    y1, y2 = int(y1), int(y2)
    name1, map1, src1, ts1 = fetch_company_year_data(co_id, y1, force=force)
    name2, map2, src2, ts2 = fetch_company_year_data(co_id, y2, force=force)
    name = name2 or name1

    avail_m2 = sorted(map2.keys())
    avail_m1 = sorted(map1.keys())
    target_m = int(ref_month) if ref_month else (avail_m2[-1] if avail_m2 else (avail_m1[-1] if avail_m1 else None))

    rows = []
    for m in range(1, 13):
        r1 = map1.get(m, {})
        r2 = map2.get(m, {})
        y1_mon, y1_cum = r1.get("month_pretax"), r1.get("cum_pretax")
        y2_mon, y2_cum = r2.get("month_pretax"), r2.get("cum_pretax")
        rows.append({
            "month": m,
            "month_label": f"{m}月",
            "y1_month_pretax": y1_mon,
            "y1_cum_pretax": y1_cum,
            "y2_month_pretax": y2_mon,
            "y2_cum_pretax": y2_cum,
            "month_yoy": calc_yoy(y2_mon, y1_mon),
            "cum_yoy": calc_yoy(y2_cum, y1_cum),
        })

    ref_row = rows[target_m - 1] if (target_m and 1 <= target_m <= 12) else None
    overall_source = "db" if (src1 == "db" and src2 == "db") else "mops"
    return {
        "co_id": str(co_id),
        "name": name,
        "label": f"{co_id} {name}".strip(),
        "y1": y1,
        "y2": y2,
        "ref_month": target_m,
        "latest_row": ref_row,
        "rows": rows,
        "source": overall_source,
        "updated_at": max(ts1, ts2),
    }


def style_worksheet(ws, pct_cols=()):
    hdr_fill = PatternFill("solid", start_color="1F497D")
    hdr_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
    font = Font(name="微軟正黑體", size=10)
    border = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
    headers = [c.value for c in ws[1]]
    for c in ws[1]:
        c.fill, c.font = hdr_fill, hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = font, border
            h = headers[c.column - 1]
            if h in pct_cols:
                c.number_format = "0.00%;[Red]-0.00%"
                c.alignment = Alignment(horizontal="right")
            elif isinstance(c.value, (int, float)):
                c.number_format = "#,##0;[Red]-#,##0"
                c.alignment = Alignment(horizontal="right")
            elif h in ("金融股", "公司代號", "公司簡稱", "報表類型", "月份"):
                c.alignment = Alignment(horizontal="center")
    for idx, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value or "")) * 1.5 for c in col)
        ws.column_dimensions[get_column_letter(idx)].width = max(min(w + 4, 30), 13)
    ws.freeze_panes = "A2"


def save_excel_file(filepath: str, payload: dict):
    ym = payload.get("ym", "11509")
    y, m = int(str(ym)[:-2]), int(str(ym)[-2:])
    y1 = int(payload.get("y1", y - 1))
    y2 = int(payload.get("y2", y))
    stocks = payload.get("stocks", [])
    companies = payload.get("companies", {})

    summary_rows = []
    for s in stocks:
        cid = s["co_id"]
        comp = companies.get(cid)
        latest = comp.get("latest_row") if comp else None
        summary_rows.append({
            "金融股": s["label"],
            "報表類型": s["report_type"],
            f"{y}/{m:02d}\n本月稅前損益(清單)": s.get("list_month_pretax"),
            f"{y}/{m:02d}\n本年累計稅前損益(清單)": s.get("list_cum_pretax"),
            f"{y1}年度\n本月稅前損益": latest.get("y1_month_pretax") if latest else None,
            f"{y1}年度\n本年累計稅前損益": latest.get("y1_cum_pretax") if latest else None,
            f"{y2}年度\n本月稅前損益": latest.get("y2_month_pretax") if latest else None,
            f"{y2}年度\n本年累計稅前損益": latest.get("y2_cum_pretax") if latest else None,
            "累計稅前 YoY": latest.get("cum_yoy") if latest else None,
        })
    df_sum = pd.DataFrame(summary_rows)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df_sum.to_excel(writer, sheet_name="自結清單", index=False)
        for cid, comp in companies.items():
            if not comp or "rows" not in comp:
                continue
            c_rows = []
            for r in comp["rows"]:
                c_rows.append({
                    "月份": r["month_label"],
                    f"{y1}年度\n本月稅前損益": r["y1_month_pretax"],
                    f"{y1}年度\n本年累計稅前損益": r["y1_cum_pretax"],
                    f"{y2}年度\n本月稅前損益": r["y2_month_pretax"],
                    f"{y2}年度\n本年累計稅前損益": r["y2_cum_pretax"],
                    "單月稅前 YoY": r["month_yoy"],
                    "累計稅前 YoY": r["cum_yoy"],
                })
            df_c = pd.DataFrame(c_rows)
            sheet_title = re.sub(r'[\\/*?:\[\]]', "_", comp.get("label", cid))[:31]
            df_c.to_excel(writer, sheet_title, index=False)

    wb = openpyxl.load_workbook(filepath)
    for ws in wb.worksheets:
        style_worksheet(ws, pct_cols=("累計稅前 YoY", "單月稅前 YoY"))
    wb.save(filepath)


# ================================================================ Windows Desktop GUI
class MopsDesktopApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("📊 MOPS 金融股月自結損益彙整系統 (Windows 桌面版)")
        self.geometry("1440x860")
        self.minsize(1120, 680)
        self.configure(bg="#0f172a")

        today = datetime.date.today()
        curr_roc = today.year - 1911
        self.default_y1 = curr_roc - 1
        self.default_y2 = curr_roc
        self.ym = "11509"
        self.list_year = curr_roc
        self.list_month = 9
        self.stocks = []
        self.selected = set()
        self.companies = {}
        self.active_tab = "SUMMARY"
        self.source_co_id = "2881"
        self.sort_col = "co_id"
        self.sort_asc = True
        self.is_updating = False
        self.stop_requested = False

        self.var_ym = tk.StringVar(value="11509")
        self.var_y1 = tk.IntVar(value=self.default_y1)
        self.var_y2 = tk.IntVar(value=self.default_y2)
        self.var_fetch_mode = tk.StringVar(value="db")
        self.var_source_co = tk.StringVar(value="2881 富邦金")
        self.stock_chk_vars = {}

        self._setup_styles()
        self._build_ui()

        # 啟動背景執行緒載入初始自結清單
        self.after(50, self.start_init_load)

    def _setup_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")

        bg_dark = "#0f172a"
        bg_card = "#1e293b"
        bg_head = "#1f497d"
        fg_text = "#f8fafc"
        fg_muted = "#94a3b8"
        border_col = "#334155"

        style.configure(".", background=bg_dark, foreground=fg_text, font=("Microsoft JhengHei UI", 10))
        style.configure("Card.TFrame", background=bg_card)
        style.configure("Dark.TFrame", background=bg_dark)
        style.configure("Card.TLabel", background=bg_card, foreground=fg_text, font=("Microsoft JhengHei UI", 10))
        style.configure("Muted.TLabel", background=bg_card, foreground=fg_muted, font=("Microsoft JhengHei UI", 9))

        style.configure(
            "Mops.Treeview",
            background=bg_card,
            fieldbackground=bg_card,
            foreground=fg_text,
            rowheight=30,
            borderwidth=0,
            font=("Consolas", 10),
        )
        style.configure(
            "Mops.Treeview.Heading",
            background=bg_head,
            foreground="#ffffff",
            font=("Microsoft JhengHei UI", 9, "bold"),
            padding=(4, 6),
            relief="flat",
        )
        style.map(
            "Mops.Treeview.Heading",
            background=[("active", "#285285")],
            foreground=[("active", "#93c5fd")],
        )
        style.map(
            "Mops.Treeview",
            background=[("selected", "#2563eb")],
            foreground=[("selected", "#ffffff")],
        )

        style.configure(
            "Mops.Horizontal.TProgressbar",
            troughcolor=bg_dark,
            background="#3b82f6",
            bordercolor=border_col,
            lightcolor="#60a5fa",
            darkcolor="#2563eb",
        )

    def _build_ui(self):
        # 1. 頂部標題列 (整合分頁切換至標題右側，如使用者箭頭所示)
        hdr = tk.Frame(self, bg="#0f172a", padx=16, pady=8)
        hdr.pack(fill="x")

        title_box = tk.Frame(hdr, bg="#0f172a")
        title_box.pack(side="left")
        tk.Label(
            title_box,
            text="📊 MOPS 金融股月自結損益彙整系統",
            bg="#0f172a",
            fg="#f8fafc",
            font=("Microsoft JhengHei UI", 14, "bold"),
        ).pack(side="left")

        # 分頁切換列直接放在標題右側 (對應截圖紅框箭頭位置)
        self.tab_bar = tk.Frame(hdr, bg="#0f172a", padx=12)
        self.tab_bar.pack(side="left", fill="x", expand=True)

        btn_box = tk.Frame(hdr, bg="#0f172a")
        btn_box.pack(side="right")
        self.btn_export_xlsx = tk.Button(
            btn_box,
            text="📥 匯出 Excel (.xlsx)",
            bg="#10b981",
            fg="#ffffff",
            activebackground="#059669",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="flat",
            padx=12,
            pady=5,
            cursor="hand2",
            command=self.export_excel,
        )
        self.btn_export_xlsx.pack(side="left", padx=4)

        tk.Button(
            btn_box,
            text="📄 匯出清單 CSV",
            bg="#1e293b",
            fg="#f8fafc",
            activebackground="#334155",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="solid",
            bd=1,
            padx=12,
            pady=5,
            cursor="hand2",
            command=self.export_csv,
        ).pack(side="left", padx=4)

        # 2. 上方雙欄控制卡片區
        ctrl_grid = tk.Frame(self, bg="#0f172a", padx=16, pady=4)
        ctrl_grid.pack(fill="x")
        ctrl_grid.columnconfigure(0, weight=5)
        ctrl_grid.columnconfigure(1, weight=6)

        # ---------------- 左側卡片 (含 3 個獨立 Tabs：設定 / 自結清單資料來源 / 個股資料來源) ----------------
        card1 = tk.Frame(ctrl_grid, bg="#1e293b", highlightbackground="#334155", highlightthickness=1, padx=12, pady=10)
        card1.grid(row=0, column=0, sticky="nsew", padx=(0, 7))

        c1_top = tk.Frame(card1, bg="#1e293b")
        c1_top.pack(fill="x", pady=(0, 8))

        self.c1_tab_btns = {}
        tabs_spec = [
            ("config", "① 清單年月 & 兩年度選擇"),
            ("srcList", "📋 自結清單 資料來源"),
            ("srcComp", "📈 個股自結損益表 資料來源"),
        ]
        for key, text in tabs_spec:
            b = tk.Button(
                c1_top,
                text=text,
                font=("Microsoft JhengHei UI", 9, "bold"),
                relief="flat",
                padx=8,
                pady=4,
                cursor="hand2",
                command=lambda k=key: self.switch_card1_tab(k),
            )
            b.pack(side="left", padx=(0, 5))
            self.c1_tab_btns[key] = b

        self.btn_reset_years = tk.Button(
            c1_top,
            text="重設預設年度",
            bg="#0f172a",
            fg="#cbd5e1",
            activebackground="#334155",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 8),
            relief="flat",
            padx=8,
            pady=3,
            cursor="hand2",
            command=self.reset_default_years,
        )
        self.btn_reset_years.pack(side="right")

        # Card 1 - Pane 1: 清單年月 & 兩年度選擇
        self.c1_pane_config = tk.Frame(card1, bg="#1e293b")
        r1 = tk.Frame(self.c1_pane_config, bg="#1e293b")
        r1.pack(fill="x", pady=4)
        tk.Label(r1, text="最新清單年月 (民國 YYYMM)：", bg="#1e293b", fg="#94a3b8", font=("Microsoft JhengHei UI", 9)).pack(side="left")
        tk.Entry(
            r1,
            textvariable=self.var_ym,
            width=9,
            bg="#0f172a",
            fg="#f8fafc",
            insertbackground="#f8fafc",
            relief="solid",
            bd=1,
            font=("Consolas", 10),
        ).pack(side="left", padx=6)
        self.btn_load_list = tk.Button(
            r1,
            text="🔄 載入清單",
            bg="#0f172a",
            fg="#f8fafc",
            activebackground="#334155",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="flat",
            padx=10,
            pady=3,
            cursor="hand2",
            command=self.on_click_load_list,
        )
        self.btn_load_list.pack(side="left", padx=4)

        r2 = tk.Frame(self.c1_pane_config, bg="#1e293b")
        r2.pack(fill="x", pady=6)
        tk.Label(r2, text="比較年度 1 (基準/上一年度)：", bg="#1e293b", fg="#94a3b8", font=("Microsoft JhengHei UI", 9)).pack(side="left")
        sp1 = tk.Spinbox(
            r2,
            from_=102,
            to=130,
            textvariable=self.var_y1,
            width=6,
            bg="#0f172a",
            fg="#f8fafc",
            buttonbackground="#334155",
            relief="solid",
            bd=1,
            font=("Consolas", 10),
            command=self.on_year_change,
        )
        sp1.pack(side="left", padx=5)
        sp1.bind("<Return>", lambda e: self.on_year_change())

        tk.Label(r2, text="vs   比較年度 2 (目標/今年度)：", bg="#1e293b", fg="#94a3b8", font=("Microsoft JhengHei UI", 9)).pack(side="left", padx=4)
        sp2 = tk.Spinbox(
            r2,
            from_=102,
            to=130,
            textvariable=self.var_y2,
            width=6,
            bg="#0f172a",
            fg="#f8fafc",
            buttonbackground="#334155",
            relief="solid",
            bd=1,
            font=("Consolas", 10),
            command=self.on_year_change,
        )
        sp2.pack(side="left", padx=5)
        sp2.bind("<Return>", lambda e: self.on_year_change())

        tk.Label(
            self.c1_pane_config,
            text="💡 預設為「上一年度」與「今年度」，可自由調整任意兩個民國年度進行 1~12 月對齊比較。",
            bg="#1e293b",
            fg="#94a3b8",
            font=("Microsoft JhengHei UI", 8),
            anchor="w",
        ).pack(fill="x", pady=(4, 0))

        # Card 1 - Pane 2: 📋 自結清單 資料來源 (URL & Payload)
        self.c1_pane_src_list = tk.Frame(card1, bg="#0f172a", padx=10, pady=8, highlightbackground="#2b4468", highlightthickness=1)
        sl_top = tk.Frame(self.c1_pane_src_list, bg="#0f172a")
        sl_top.pack(fill="x", pady=(0, 6))
        tk.Label(sl_top, text="📋 自結清單 (彙總表) 資料來源", bg="#0f172a", fg="#93c5fd", font=("Microsoft JhengHei UI", 9, "bold")).pack(side="left")

        tk.Button(
            sl_top,
            text="🌐 MOPS 查詢頁 (t138sb01) ↗",
            bg="#172554",
            fg="#38bdf8",
            activebackground="#1e3a8a",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=7,
            pady=2,
            cursor="hand2",
            command=lambda: webbrowser.open("https://mops.twse.com.tw/mops/#/web/t138sb01"),
        ).pack(side="right", padx=3)

        self.btn_open_list_mops = tk.Button(
            sl_top,
            text="🔗 點此直接開啟 11509 原始報表 ↗",
            bg="#172554",
            fg="#38bdf8",
            activebackground="#1e3a8a",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=7,
            pady=2,
            cursor="hand2",
            command=self.open_mops_list_report,
        )
        self.btn_open_list_mops.pack(side="right", padx=3)

        self.txt_list_payload = tk.Label(
            self.c1_pane_src_list,
            text="",
            bg="#0b1120",
            fg="#7dd3fc",
            font=("Consolas", 8),
            justify="left",
            anchor="w",
            padx=8,
            pady=6,
            cursor="hand2",
            wraplength=540,
        )
        self.txt_list_payload.pack(fill="x")
        self.txt_list_payload.bind("<Button-1>", lambda e: self.open_mops_list_report())

        # Card 1 - Pane 3: 📈 個股自結損益表 資料來源 (URL & Payload)
        self.c1_pane_src_comp = tk.Frame(card1, bg="#0f172a", padx=10, pady=8, highlightbackground="#2b4468", highlightthickness=1)
        sc_top = tk.Frame(self.c1_pane_src_comp, bg="#0f172a")
        sc_top.pack(fill="x", pady=(0, 6))
        tk.Label(sc_top, text="📈 個股自結表來源", bg="#0f172a", fg="#93c5fd", font=("Microsoft JhengHei UI", 9, "bold")).pack(side="left")

        self.cmb_source_co = ttk.Combobox(
            sc_top,
            textvariable=self.var_source_co,
            width=13,
            state="readonly",
            font=("Microsoft JhengHei UI", 8),
        )
        self.cmb_source_co.pack(side="left", padx=6)
        self.cmb_source_co.bind("<<ComboboxSelected>>", self.on_source_combo_selected)

        tk.Button(
            sc_top,
            text="🌐 查詢頁 (t138sb02_q1) ↗",
            bg="#172554",
            fg="#38bdf8",
            activebackground="#1e3a8a",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=6,
            pady=2,
            cursor="hand2",
            command=lambda: webbrowser.open("https://mops.twse.com.tw/mops/#/web/t138sb02_q1"),
        ).pack(side="right", padx=2)

        self.btn_open_comp_y2 = tk.Button(
            sc_top,
            text="🔗 開啟 115年報表 ↗",
            bg="#172554",
            fg="#38bdf8",
            activebackground="#1e3a8a",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=6,
            pady=2,
            cursor="hand2",
            command=lambda: self.open_mops_company_report(self.source_co_id, self.var_y2.get()),
        )
        self.btn_open_comp_y2.pack(side="right", padx=2)

        self.btn_open_comp_y1 = tk.Button(
            sc_top,
            text="🔗 開啟 114年報表 ↗",
            bg="#172554",
            fg="#38bdf8",
            activebackground="#1e3a8a",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=6,
            pady=2,
            cursor="hand2",
            command=lambda: self.open_mops_company_report(self.source_co_id, self.var_y1.get()),
        )
        self.btn_open_comp_y1.pack(side="right", padx=2)

        self.txt_comp_payload = tk.Label(
            self.c1_pane_src_comp,
            text="",
            bg="#0b1120",
            fg="#7dd3fc",
            font=("Consolas", 8),
            justify="left",
            anchor="w",
            padx=8,
            pady=5,
            cursor="hand2",
            wraplength=540,
        )
        self.txt_comp_payload.pack(fill="x")
        self.txt_comp_payload.bind("<Button-1>", lambda e: self.open_mops_company_report(self.source_co_id, self.var_y2.get()))

        self.switch_card1_tab("config")

        # ---------------- 右側卡片 (② 選擇要更新月自結損益的金融股) ----------------
        card2 = tk.Frame(ctrl_grid, bg="#1e293b", highlightbackground="#334155", highlightthickness=1, padx=12, pady=10)
        card2.grid(row=0, column=1, sticky="nsew", padx=(7, 0))

        c2_top = tk.Frame(card2, bg="#1e293b")
        c2_top.pack(fill="x", pady=(0, 6))
        self.lbl_select_title = tk.Label(
            c2_top,
            text="② 選擇要更新月自結損益的金融股 (0 / 0 檔)",
            bg="#1e293b",
            fg="#cbd5e1",
            font=("Microsoft JhengHei UI", 9, "bold"),
        )
        self.lbl_select_title.pack(side="left")

        preset_bar = tk.Frame(c2_top, bg="#1e293b")
        preset_bar.pack(side="right")
        for lbl, cmd in [
            ("預設五檔", lambda: self.select_preset(["2881", "2882", "2891", "2884", "2885"])),
            ("前三檔", lambda: self.select_top_n(3)),
            ("全選 (ALL)", lambda: self.select_all(True)),
            ("清除", lambda: self.select_all(False)),
        ]:
            tk.Button(
                preset_bar,
                text=lbl,
                bg="#0f172a",
                fg="#cbd5e1",
                activebackground="#334155",
                activeforeground="#ffffff",
                font=("Microsoft JhengHei UI", 8),
                relief="flat",
                padx=7,
                pady=2,
                cursor="hand2",
                command=cmd,
            ).pack(side="left", padx=2)

        chips_outer = tk.Frame(card2, bg="#0f172a", highlightbackground="#334155", highlightthickness=1)
        chips_outer.pack(fill="x", pady=4)
        self.chips_canvas = tk.Canvas(chips_outer, bg="#0f172a", height=74, highlightthickness=0)
        chips_scroll = ttk.Scrollbar(chips_outer, orient="vertical", command=self.chips_canvas.yview)
        self.chips_inner = tk.Frame(self.chips_canvas, bg="#0f172a", padx=6, pady=4)
        self.chips_inner.bind(
            "<Configure>",
            lambda e: self.chips_canvas.configure(scrollregion=self.chips_canvas.bbox("all")),
        )
        self.chips_canvas.create_window((0, 0), window=self.chips_inner, anchor="nw")
        self.chips_canvas.configure(yscrollcommand=chips_scroll.set)
        self.chips_canvas.pack(side="left", fill="both", expand=True)
        chips_scroll.pack(side="right", fill="y")

        c2_bot = tk.Frame(card2, bg="#1e293b")
        c2_bot.pack(fill="x", pady=(6, 0))

        act_left = tk.Frame(c2_bot, bg="#1e293b")
        act_left.pack(side="left")

        self.btn_update_sel = tk.Button(
            act_left,
            text="🚀 更新選取個股 (0 檔)",
            bg="#3b82f6",
            fg="#ffffff",
            activebackground="#2563eb",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="flat",
            padx=10,
            pady=4,
            cursor="hand2",
            command=self.update_selected_stocks,
        )
        self.btn_update_sel.pack(side="left", padx=(0, 6))

        self.btn_update_all = tk.Button(
            act_left,
            text="⚡ 一鍵更新全部 (ALL)",
            bg="#f59e0b",
            fg="#111827",
            activebackground="#d97706",
            activeforeground="#111827",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="flat",
            padx=10,
            pady=4,
            cursor="hand2",
            command=self.update_all_stocks,
        )
        self.btn_update_all.pack(side="left", padx=(0, 6))

        self.btn_stop = tk.Button(
            act_left,
            text="⏹ 停止",
            bg="#ef4444",
            fg="#ffffff",
            activebackground="#dc2626",
            activeforeground="#ffffff",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="flat",
            padx=10,
            pady=4,
            cursor="hand2",
            command=self.stop_updating,
        )

        # 模式切換：左邊「🔄 強制重抓」、右邊「🗄️ 資料庫優先」
        mode_frame = tk.Frame(c2_bot, bg="#0f172a", padx=3, pady=3, highlightbackground="#334155", highlightthickness=1)
        mode_frame.pack(side="right")

        self.btn_mode_force = tk.Button(
            mode_frame,
            text="🔄 強制重抓",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=8,
            pady=3,
            cursor="hand2",
            command=lambda: self.set_fetch_mode("force"),
        )
        self.btn_mode_force.pack(side="left", padx=2)

        self.btn_mode_db = tk.Button(
            mode_frame,
            text="🗄️ 資料庫優先",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=8,
            pady=3,
            cursor="hand2",
            command=lambda: self.set_fetch_mode("db"),
        )
        self.btn_mode_db.pack(side="left", padx=2)
        self._refresh_mode_buttons()

        # 3. 狀態與進度列
        status_bar = tk.Frame(self, bg="#1e293b", highlightbackground="#334155", highlightthickness=1, padx=14, pady=8)
        status_bar.pack(fill="x", padx=16, pady=6)

        st_left = tk.Frame(status_bar, bg="#1e293b")
        st_left.pack(side="left", fill="x", expand=True)

        st_top = tk.Frame(st_left, bg="#1e293b")
        st_top.pack(fill="x")
        self.lbl_status = tk.Label(
            st_top,
            text="正在初始化並讀取最新月自結清單...",
            bg="#1e293b",
            fg="#f8fafc",
            font=("Microsoft JhengHei UI", 9),
            anchor="w",
        )
        self.lbl_status.pack(side="left")
        self.lbl_progress_cnt = tk.Label(
            st_top,
            text="",
            bg="#1e293b",
            fg="#94a3b8",
            font=("Consolas", 9),
        )
        self.lbl_progress_cnt.pack(side="right", padx=10)

        self.progress = ttk.Progressbar(st_left, style="Mops.Horizontal.TProgressbar", orient="horizontal", mode="determinate")
        self.progress.pack(fill="x", pady=(4, 0))

        tk.Label(
            status_bar,
            text="單位：新台幣仟元 ｜ 點擊表頭可排序 ｜ 雙擊個股檢視 1~12 月對齊表 ｜ 右鍵點擊個股可開啟 MOPS 原始報表",
            bg="#1e293b",
            fg="#94a3b8",
            font=("Microsoft JhengHei UI", 8),
        ).pack(side="right", padx=(14, 0))

        # 4. 主表格容器 (自結清單視圖 vs 個股 1~12 月明細視圖)
        self.main_area = tk.Frame(self, bg="#0f172a", padx=16, pady=4)
        self.main_area.pack(fill="both", expand=True)

        # ---- 視圖 A: 自結清單 Treeview ----
        self.view_summary = tk.Frame(self.main_area, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)
        self.view_summary.pack(fill="both", expand=True)

        cols = (
            "chk",
            "co_id",
            "report_type",
            "list_month_pretax",
            "list_cum_pretax",
            "y1_month_pretax",
            "y1_cum_pretax",
            "y2_month_pretax",
            "y2_cum_pretax",
            "cum_yoy",
            "status",
        )
        self.tree_sum = ttk.Treeview(self.view_summary, columns=cols, show="headings", style="Mops.Treeview", selectmode="browse")
        vsb = ttk.Scrollbar(self.view_summary, orient="vertical", command=self.tree_sum.yview)
        hsb = ttk.Scrollbar(self.view_summary, orient="horizontal", command=self.tree_sum.xview)
        self.tree_sum.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree_sum.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        col_widths = {
            "chk": (48, "center"),
            "co_id": (140, "w"),
            "report_type": (76, "center"),
            "list_month_pretax": (145, "e"),
            "list_cum_pretax": (160, "e"),
            "y1_month_pretax": (145, "e"),
            "y1_cum_pretax": (160, "e"),
            "y2_month_pretax": (145, "e"),
            "y2_cum_pretax": (160, "e"),
            "cum_yoy": (105, "e"),
            "status": (185, "center"),
        }
        for c, (w, anc) in col_widths.items():
            self.tree_sum.column(c, width=w, minwidth=45, anchor=anc)

        self.tree_sum.tag_configure("updated", background="#162a3a", foreground="#f8fafc")
        self.tree_sum.tag_configure("normal", background="#1e293b", foreground="#f8fafc")

        self.tree_sum.bind("<ButtonRelease-1>", self.on_summary_tree_click)
        self.tree_sum.bind("<Double-1>", self.on_summary_tree_double_click)
        self.tree_sum.bind("<Button-3>", self.on_summary_tree_right_click)

        self.ctx_menu = tk.Menu(self, tearoff=0, bg="#1e293b", fg="#f8fafc", activebackground="#2563eb", activeforeground="#ffffff")

        # ---- 視圖 B: 個股 1~12 月明細視圖 ----
        self.view_company = tk.Frame(self.main_area, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)

        vc_hdr = tk.Frame(self.view_company, bg="#1e293b", padx=14, pady=10)
        vc_hdr.pack(fill="x")

        vc_left = tk.Frame(vc_hdr, bg="#1e293b")
        vc_left.pack(side="left", fill="x", expand=True)

        vc_title_row = tk.Frame(vc_left, bg="#1e293b")
        vc_title_row.pack(fill="x")
        self.lbl_comp_title = tk.Label(vc_title_row, text="", bg="#1e293b", fg="#93c5fd", font=("Microsoft JhengHei UI", 11, "bold"))
        self.lbl_comp_title.pack(side="left")
        self.lbl_comp_sub = tk.Label(vc_title_row, text="", bg="#1e293b", fg="#94a3b8", font=("Microsoft JhengHei UI", 9))
        self.lbl_comp_sub.pack(side="left", padx=12)

        vc_src_row = tk.Frame(vc_left, bg="#0b1120", padx=8, pady=5)
        vc_src_row.pack(fill="x", pady=(6, 0))
        self.btn_vc_open_y1 = tk.Button(
            vc_src_row,
            text="🔗 點擊開啟 114年度 MOPS 原始報表 ↗",
            bg="#172554",
            fg="#38bdf8",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=7,
            pady=2,
            cursor="hand2",
        )
        self.btn_vc_open_y1.pack(side="left", padx=(0, 6))
        self.btn_vc_open_y2 = tk.Button(
            vc_src_row,
            text="🔗 點擊開啟 115年度 MOPS 原始報表 ↗",
            bg="#172554",
            fg="#38bdf8",
            font=("Microsoft JhengHei UI", 8, "bold"),
            relief="flat",
            padx=7,
            pady=2,
            cursor="hand2",
        )
        self.btn_vc_open_y2.pack(side="left", padx=(0, 8))
        self.lbl_vc_payload = tk.Label(vc_src_row, text="", bg="#0b1120", fg="#cbd5e1", font=("Consolas", 8), anchor="w")
        self.lbl_vc_payload.pack(side="left", fill="x", expand=True)

        tk.Button(
            vc_hdr,
            text="← 返回自結清單",
            bg="#0f172a",
            fg="#f8fafc",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="flat",
            padx=12,
            pady=5,
            cursor="hand2",
            command=lambda: self.switch_tab("SUMMARY"),
        ).pack(side="right", padx=4)

        comp_cols = ("month", "y1_mon", "y1_cum", "y2_mon", "y2_cum", "mon_yoy", "cum_yoy")
        self.tree_comp = ttk.Treeview(self.view_company, columns=comp_cols, show="headings", style="Mops.Treeview", selectmode="browse")
        vc_vsb = ttk.Scrollbar(self.view_company, orient="vertical", command=self.tree_comp.yview)
        self.tree_comp.configure(yscrollcommand=vc_vsb.set)
        self.tree_comp.pack(side="left", fill="both", expand=True)
        vc_vsb.pack(side="right", fill="y")

        for c, (w, anc) in [
            ("month", (90, "center")),
            ("y1_mon", (180, "e")),
            ("y1_cum", (190, "e")),
            ("y2_mon", (180, "e")),
            ("y2_cum", (190, "e")),
            ("mon_yoy", (130, "e")),
            ("cum_yoy", (130, "e")),
        ]:
            self.tree_comp.column(c, width=w, anchor=anc)
        self.tree_comp.tag_configure("ref_month", background="#1e3a5f", foreground="#93c5fd")

        self.render_header_labels()
        self.render_tabs()

    # ================================================================ Card 1 Tabs & 資料來源
    def switch_card1_tab(self, key: str):
        for k, btn in self.c1_tab_btns.items():
            if k == key:
                btn.configure(bg="#1d4ed8", fg="#ffffff")
            else:
                btn.configure(bg="#0f172a", fg="#94a3b8")

        self.c1_pane_config.pack_forget()
        self.c1_pane_src_list.pack_forget()
        self.c1_pane_src_comp.pack_forget()

        if key == "config":
            self.c1_pane_config.pack(fill="both", expand=True)
            self.btn_reset_years.pack(side="right")
        elif key == "srcList":
            self.c1_pane_src_list.pack(fill="both", expand=True)
            self.btn_reset_years.pack_forget()
        elif key == "srcComp":
            self.c1_pane_src_comp.pack(fill="both", expand=True)
            self.btn_reset_years.pack_forget()
        self.render_source_info()

    def on_source_combo_selected(self, _event=None):
        val = self.var_source_co.get().strip()
        co_id = val.split()[0] if val else "2881"
        self.set_source_co_id(co_id)

    def set_source_co_id(self, co_id: str):
        self.source_co_id = str(co_id)
        for s in self.stocks:
            if s["co_id"] == self.source_co_id:
                self.var_source_co.set(s["label"])
                break
        self.render_source_info()

    def render_source_info(self):
        ym = self.var_ym.get().strip() or self.ym or "11509"
        y1 = self.var_y1.get()
        y2 = self.var_y2.get()
        co_id = self.active_tab if self.active_tab != "SUMMARY" else (self.source_co_id or "2881")

        list_payload = json.dumps(
            {
                "apiName": "ajax_t138sb01",
                "parameters": {"CK2": "1", "YM1": str(ym), "CK1": "1", "encodeURIComponent": 1, "firstin": 1, "step": 1, "first": 1},
            },
            ensure_ascii=False,
        )
        comp_q1_payload = json.dumps(
            {
                "apiName": "ajax_t138sb02_q1",
                "parameters": {"co_id": str(co_id), "year": str(y2), "encodeURIComponent": 1, "step": 1, "firstin": 1, "off": 1},
            },
            ensure_ascii=False,
        )
        comp_data_payload = json.dumps(
            {
                "apiName": "ajax_t138sb02",
                "parameters": {"encodeURIComponent": 1, "step": 1, "firstin": 1, "off": 1, "CK2": "1", "BK1": "2", "TYPEK": "sii", "YEAR": str(y2), "COMP": str(co_id)},
            },
            ensure_ascii=False,
        )

        self.btn_open_list_mops.configure(text=f"🔗 點此直接開啟 {ym} 原始報表 ↗")
        self.txt_list_payload.configure(
            text=f"POST URL: {API_URL}\nPayload: {list_payload}   (點擊此處直接開啟報表)"
        )

        self.btn_open_comp_y1.configure(text=f"🔗 開啟 {co_id} ({y1}) 報表 ↗")
        self.btn_open_comp_y2.configure(text=f"🔗 開啟 {co_id} ({y2}) 報表 ↗")
        self.txt_comp_payload.configure(
            text=f"POST URL: {API_URL}\nPayload (查詢): {comp_q1_payload}\nPayload (明細): {comp_data_payload}"
        )

    def open_mops_list_report(self):
        ym = self.var_ym.get().strip() or self.ym or "11509"
        self.set_status(f"正在向 MOPS 取得 {ym} 自結清單報表簽章網址並開啟瀏覽器...")

        def _worker():
            try:
                url = resolve_mops_url(
                    "ajax_t138sb01",
                    {"CK2": "1", "YM1": str(ym), "CK1": "1", "encodeURIComponent": 1, "firstin": 1, "step": 1, "first": 1},
                )
                webbrowser.open(url)
                self.after(0, lambda: self.set_status(f"🌐 已在瀏覽器開啟 MOPS {ym} 自結清單原始報表。"))
            except Exception as e:
                err_msg = str(e)
                self.after(0, lambda m=err_msg: messagebox.showerror("開啟 MOPS 報表失敗", m))

        threading.Thread(target=_worker, daemon=True).start()

    def open_mops_company_report(self, co_id: str, year: int):
        co_id = str(co_id).strip()
        year = str(year).strip()
        self.set_status(f"正在向 MOPS 取得 {co_id} ({year}年度) 個股自結損益表網址並開啟瀏覽器...")

        def _worker():
            try:
                url = resolve_mops_url(
                    "ajax_t138sb02",
                    {
                        "encodeURIComponent": 1,
                        "step": 1,
                        "firstin": 1,
                        "off": 1,
                        "CK2": "1",
                        "BK1": "2",
                        "TYPEK": "sii",
                        "YEAR": year,
                        "COMP": co_id,
                    },
                )
                webbrowser.open(url)
                self.after(0, lambda: self.set_status(f"🌐 已在瀏覽器開啟 MOPS {co_id} ({year}年度) 個股自結損益表。"))
            except Exception as e:
                err_msg = str(e)
                self.after(0, lambda m=err_msg: messagebox.showerror("開啟 MOPS 報表失敗", m))

        threading.Thread(target=_worker, daemon=True).start()

    # ================================================================ 模式切換 & 狀態列
    def set_fetch_mode(self, mode: str):
        self.var_fetch_mode.set(mode)
        self._refresh_mode_buttons()
        if mode == "force":
            self.set_status("🔄 已切換為「強制重抓」模式：每次執行皆重新向 MOPS 抓取並同步更新至 SQLite3 資料庫。")
        else:
            self.set_status("🗄️ 已切換為「資料庫優先」模式：若 SQLite3 資料庫已有資料則直接採用，無資料時才向 MOPS 抓取並寫入資料庫。")

    def _refresh_mode_buttons(self):
        is_force = self.var_fetch_mode.get() == "force"
        if is_force:
            self.btn_mode_force.configure(bg="#b45309", fg="#fef3c7")
            self.btn_mode_db.configure(bg="#0f172a", fg="#94a3b8")
        else:
            self.btn_mode_force.configure(bg="#0f172a", fg="#94a3b8")
            self.btn_mode_db.configure(bg="#065f46", fg="#6ee7b7")

    def set_status(self, msg: str, done: int = 0, total: int = 0):
        self.lbl_status.configure(text=msg)
        pct = round((done / total) * 100) if total > 0 else 0
        self.progress["value"] = pct
        self.lbl_progress_cnt.configure(text=f"{done} / {total} ({pct}%)" if total > 0 else "")

    # ================================================================ 資料載入與事件
    def start_init_load(self):
        self.set_status("正在讀取最新月自結損益清單...")

        def _worker():
            try:
                list_data = find_latest_ym(market="1")
                self.after(0, lambda d=list_data: self._on_init_loaded(d))
            except Exception as e:
                err_msg = str(e)
                self.after(0, lambda m=err_msg: self.set_status(f"❌ 初始化錯誤: {m}"))

        threading.Thread(target=_worker, daemon=True).start()

    def _on_init_loaded(self, list_data: dict):
        self.var_ym.set(list_data["ym"])
        self.apply_list_data(list_data)
        self.set_status(
            f"已載入 {self.list_year}年{self.list_month}月 金融股自結清單（共 {len(self.stocks)} 檔）。請勾選個股後點擊「🚀 更新選取個股」或「⚡ 一鍵更新全部 (ALL)」。"
        )

    def on_click_load_list(self):
        ym = self.var_ym.get().strip()
        if not ym:
            return
        force = self.var_fetch_mode.get() == "force"
        self.btn_load_list.configure(state="disabled")
        self.set_status(f"正在讀取 {ym} 月自結損益清單 ({'強制重抓' if force else '資料庫優先'})...")

        def _worker():
            try:
                list_data = fetch_list_data(ym, market="1", force=force)
                self.after(0, lambda d=list_data: self._on_list_reloaded(d))
            except Exception as e:
                err_msg = str(e)
                self.after(0, lambda m=err_msg: self.set_status(f"❌ 讀取清單錯誤: {m}"))
            finally:
                self.after(0, lambda: self.btn_load_list.configure(state="normal"))

        threading.Thread(target=_worker, daemon=True).start()

    def _on_list_reloaded(self, list_data: dict):
        self.apply_list_data(list_data)
        src_txt = "🗄️ 來自 SQLite 資料庫" if list_data.get("source") == "db" else "🌐 已從 MOPS 抓取並寫入資料庫"
        self.set_status(f"已載入 {self.ym} 清單（{src_txt}），共擷取 {len(self.stocks)} 檔金融股。")

    def apply_list_data(self, list_data: dict):
        self.ym = list_data["ym"]
        self.list_year = list_data["year"]
        self.list_month = list_data["month"]
        self.stocks = list_data.get("stocks", [])

        valid_ids = {s["co_id"] for s in self.stocks}
        self.selected = {cid for cid in self.selected if cid in valid_ids}

        labels = [s["label"] for s in self.stocks]
        self.cmb_source_co["values"] = labels
        if self.stocks:
            if self.source_co_id not in valid_ids:
                self.source_co_id = "2881" if "2881" in valid_ids else self.stocks[0]["co_id"]
            self.set_source_co_id(self.source_co_id)

        self.render_header_labels()
        self.render_chips()
        self.render_summary_table()
        self.render_tabs()

    def reset_default_years(self):
        self.var_y1.set(self.default_y1)
        self.var_y2.set(self.default_y2)
        self.on_year_change()

    def on_year_change(self):
        self.companies = {}
        self.render_header_labels()
        self.render_chips()
        self.render_summary_table()
        self.render_tabs()
        y1, y2 = self.var_y1.get(), self.var_y2.get()
        self.set_status(f"比較年度已切換為 {y1}年度 vs {y2}年度，請點擊「更新選取個股」或「一鍵更新全部 (ALL)」抓取對應年度資料。")

    # ================================================================ 排序與表格渲染
    def sort_icon(self, col: str) -> str:
        if self.sort_col != col:
            return " ⇅"
        return " ▲" if self.sort_asc else " ▼"

    def sort_table(self, col: str):
        text_cols = {"co_id", "report_type"}
        if self.sort_col == col:
            self.sort_asc = not self.sort_asc
        else:
            self.sort_col = col
            self.sort_asc = col in text_cols
        self.render_header_labels()
        self.render_summary_table()

    def get_sort_value(self, s: dict, col: str):
        comp = self.companies.get(s["co_id"])
        latest = comp.get("latest_row") if comp else None
        if col == "co_id":
            return s["co_id"]
        if col == "report_type":
            return s.get("report_type") or ""
        if col == "list_month_pretax":
            return s.get("list_month_pretax")
        if col == "list_cum_pretax":
            return s.get("list_cum_pretax")
        if col == "y1_month_pretax":
            return latest.get("y1_month_pretax") if latest else None
        if col == "y1_cum_pretax":
            return latest.get("y1_cum_pretax") if latest else None
        if col == "y2_month_pretax":
            return latest.get("y2_month_pretax") if latest else None
        if col == "y2_cum_pretax":
            return latest.get("y2_cum_pretax") if latest else None
        if col == "cum_yoy":
            return latest.get("cum_yoy") if latest else None
        if col == "status":
            return 1 if comp else 0
        return s["co_id"]

    def get_sorted_stocks(self):
        col = self.sort_col
        asc = self.sort_asc
        non_null = []
        nulls = []
        for s in self.stocks:
            v = self.get_sort_value(s, col)
            if v is None or v == "":
                nulls.append(s)
            else:
                non_null.append((v, s["co_id"], s))
        non_null.sort(key=lambda item: (item[0], item[1]), reverse=not asc)
        nulls.sort(key=lambda item: item["co_id"])
        return [item[2] for item in non_null] + nulls

    def render_header_labels(self):
        y1 = self.var_y1.get()
        y2 = self.var_y2.get()
        ym_lbl = f"{self.list_year}/{self.list_month:02d}"
        m_lbl = f"({self.list_month}月)"

        all_checked = len(self.stocks) > 0 and len(self.selected) == len(self.stocks)
        self.tree_sum.heading("chk", text="☑" if all_checked else "☐", command=lambda: self.select_all(not all_checked))

        headings = [
            ("co_id", "金融股 (代號 名稱)"),
            ("report_type", "報表類型"),
            ("list_month_pretax", f"{ym_lbl}\n本月稅前損益(清單)"),
            ("list_cum_pretax", f"{ym_lbl}\n本年累計稅前損益(清單)"),
            ("y1_month_pretax", f"{y1}年度\n本月稅前損益{m_lbl}"),
            ("y1_cum_pretax", f"{y1}年度\n本年累計稅前損益{m_lbl}"),
            ("y2_month_pretax", f"{y2}年度\n本月稅前損益{m_lbl}"),
            ("y2_cum_pretax", f"{y2}年度\n本年累計稅前損益{m_lbl}"),
            ("cum_yoy", "累計稅前\nYoY"),
            ("status", "狀態 / 操作"),
        ]
        for col, text in headings:
            self.tree_sum.heading(col, text=f"{text}{self.sort_icon(col)}", command=lambda c=col: self.sort_table(c))

        self.tree_comp.heading("month", text="月份")
        self.tree_comp.heading("y1_mon", text=f"{y1}年度\n本月稅前損益")
        self.tree_comp.heading("y1_cum", text=f"{y1}年度\n本年累計稅前損益")
        self.tree_comp.heading("y2_mon", text=f"{y2}年度\n本月稅前損益")
        self.tree_comp.heading("y2_cum", text=f"{y2}年度\n本年累計稅前損益")
        self.tree_comp.heading("mon_yoy", text="單月稅前\nYoY")
        self.tree_comp.heading("cum_yoy", text="累計稅前\nYoY")

        self.render_source_info()

    def render_chips(self):
        for w in self.chips_inner.winfo_children():
            w.destroy()
        self.stock_chk_vars.clear()

        cols_per_row = 5
        for idx, s in enumerate(self.stocks):
            cid = s["co_id"]
            var = tk.BooleanVar(value=(cid in self.selected))
            self.stock_chk_vars[cid] = var
            is_upd = cid in self.companies
            fg_col = "#34d399" if is_upd else "#e2e8f0"
            cb = tk.Checkbutton(
                self.chips_inner,
                text=s["label"],
                variable=var,
                bg="#0f172a",
                fg=fg_col,
                selectcolor="#1e293b",
                activebackground="#0f172a",
                activeforeground="#ffffff",
                font=("Microsoft JhengHei UI", 9),
                cursor="hand2",
                command=lambda c=cid: self.toggle_stock(c),
            )
            cb.grid(row=idx // cols_per_row, column=idx % cols_per_row, sticky="w", padx=6, pady=2)

        sel_cnt = len(self.selected)
        self.lbl_select_title.configure(text=f"② 選擇要更新月自結損益的金融股 ({sel_cnt} / {len(self.stocks)} 檔)")
        self.btn_update_sel.configure(text=f"🚀 更新選取個股 ({sel_cnt} 檔)")

    def toggle_stock(self, co_id: str):
        if co_id in self.selected:
            self.selected.remove(co_id)
        else:
            self.selected.add(co_id)
        self.set_source_co_id(co_id)
        self.render_header_labels()
        self.render_chips()
        self.render_summary_table()

    def select_all(self, flag: bool):
        if flag:
            self.selected = {s["co_id"] for s in self.stocks}
        else:
            self.selected.clear()
        self.render_header_labels()
        self.render_chips()
        self.render_summary_table()

    def select_preset(self, codes):
        valid = {s["co_id"] for s in self.stocks}
        self.selected = {c for c in codes if c in valid}
        self.render_header_labels()
        self.render_chips()
        self.render_summary_table()

    def select_top_n(self, n: int):
        self.selected = {s["co_id"] for s in self.stocks[:n]}
        self.render_header_labels()
        self.render_chips()
        self.render_summary_table()

    @staticmethod
    def fmt_num(v) -> str:
        if v is None or v == "":
            return "—"
        try:
            return f"{int(round(float(v))):,}"
        except Exception:
            return "—"

    @staticmethod
    def fmt_yoy(v) -> str:
        if v is None:
            return "—"
        try:
            pct = float(v) * 100
            return f"+{pct:.2f}%" if pct > 0 else f"{pct:.2f}%"
        except Exception:
            return "—"

    def render_summary_table(self, loading_co_id: str = None):
        for item in self.tree_sum.get_children():
            self.tree_sum.delete(item)

        sorted_list = self.get_sorted_stocks()
        for s in sorted_list:
            cid = s["co_id"]
            is_sel = cid in self.selected
            comp = self.companies.get(cid)
            latest = comp.get("latest_row") if comp else None

            if loading_co_id == cid:
                status_txt = "⏳ 讀取中..."
            elif comp:
                status_txt = "🗄️ 資料庫 (點此看明細)" if comp.get("source") == "db" else "🌐 已抓更新DB (點此看明細)"
            else:
                status_txt = "🔄 點擊單獨更新"

            name_txt = f"{s['label']} 🔍" if comp else s["label"]
            values = (
                "☑" if is_sel else "☐",
                name_txt,
                s["report_type"],
                self.fmt_num(s.get("list_month_pretax")),
                self.fmt_num(s.get("list_cum_pretax")),
                self.fmt_num(latest.get("y1_month_pretax")) if latest else "—",
                self.fmt_num(latest.get("y1_cum_pretax")) if latest else "—",
                self.fmt_num(latest.get("y2_month_pretax")) if latest else "—",
                self.fmt_num(latest.get("y2_cum_pretax")) if latest else "—",
                self.fmt_yoy(latest.get("cum_yoy")) if latest else "—",
                status_txt,
            )
            self.tree_sum.insert("", "end", iid=cid, values=values, tags=("updated" if comp else "normal",))

    def on_summary_tree_click(self, event):
        region = self.tree_sum.identify("region", event.x, event.y)
        if region != "cell":
            return
        row_id = self.tree_sum.identify_row(event.y)
        col_id = self.tree_sum.identify_column(event.x)
        if not row_id:
            return
        self.set_source_co_id(row_id)
        if col_id == "#1":
            self.toggle_stock(row_id)
        elif col_id == "#11":
            if row_id in self.companies:
                self.switch_tab(row_id)
            else:
                self.update_single_stock(row_id)

    def on_summary_tree_double_click(self, event):
        row_id = self.tree_sum.identify_row(event.y)
        if not row_id:
            return
        if row_id in self.companies:
            self.switch_tab(row_id)
        else:
            self.update_single_stock(row_id)

    def on_summary_tree_right_click(self, event):
        row_id = self.tree_sum.identify_row(event.y)
        if not row_id:
            return
        self.tree_sum.selection_set(row_id)
        self.set_source_co_id(row_id)
        y1, y2 = self.var_y1.get(), self.var_y2.get()
        stock_obj = next((s for s in self.stocks if s["co_id"] == row_id), None)
        label = stock_obj["label"] if stock_obj else row_id

        self.ctx_menu.delete(0, "end")
        self.ctx_menu.add_command(label=f"🚀 單獨更新 {label}", command=lambda: self.update_single_stock(row_id))
        if row_id in self.companies:
            self.ctx_menu.add_command(label=f"📈 檢視 {label} 1~12月完整對齊表", command=lambda: self.switch_tab(row_id))
        self.ctx_menu.add_separator()
        self.ctx_menu.add_command(
            label=f"🔗 開啟 MOPS {label} ({y1}年度) 原始報表 ↗",
            command=lambda: self.open_mops_company_report(row_id, y1),
        )
        self.ctx_menu.add_command(
            label=f"🔗 開啟 MOPS {label} ({y2}年度) 原始報表 ↗",
            command=lambda: self.open_mops_company_report(row_id, y2),
        )
        self.ctx_menu.tk_popup(event.x_root, event.y_root)

    # ================================================================ 頂部分頁切換與刪除
    def render_tabs(self):
        for w in self.tab_bar.winfo_children():
            w.destroy()

        updated_stocks = [s for s in self.stocks if s["co_id"] in self.companies]
        is_sum = self.active_tab == "SUMMARY"
        ym_txt = f"{self.list_year}年{self.list_month}月 " if self.stocks else ""
        btn_sum = tk.Button(
            self.tab_bar,
            text=f"📋 {ym_txt}自結清單 ({len(updated_stocks)}/{len(self.stocks)} 已更新)",
            bg="#3b82f6" if is_sum else "#1e293b",
            fg="#ffffff" if is_sum else "#94a3b8",
            font=("Microsoft JhengHei UI", 9, "bold"),
            relief="flat",
            padx=10,
            pady=4,
            cursor="hand2",
            command=lambda: self.switch_tab("SUMMARY"),
        )
        btn_sum.pack(side="left", padx=(0, 5), pady=2)

        for s in updated_stocks:
            cid = s["co_id"]
            is_act = self.active_tab == cid
            f = tk.Frame(self.tab_bar, bg="#3b82f6" if is_act else "#1e293b")
            f.pack(side="left", padx=(0, 5), pady=2)
            tk.Button(
                f,
                text=f"📈 {s['label']}",
                bg="#3b82f6" if is_act else "#1e293b",
                fg="#ffffff" if is_act else "#cbd5e1",
                font=("Microsoft JhengHei UI", 9, "bold"),
                relief="flat",
                padx=7,
                pady=3,
                cursor="hand2",
                command=lambda c=cid: self.switch_tab(c),
            ).pack(side="left")
            tk.Button(
                f,
                text="×",
                bg="#3b82f6" if is_act else "#1e293b",
                fg="#fca5a5" if is_act else "#94a3b8",
                activebackground="#ef4444",
                activeforeground="#ffffff",
                font=("Microsoft JhengHei UI", 9, "bold"),
                relief="flat",
                padx=5,
                pady=3,
                cursor="hand2",
                command=lambda c=cid: self.remove_company_tab(c),
            ).pack(side="left")

        if updated_stocks:
            tk.Button(
                self.tab_bar,
                text="🗑️ 清除所有個股分頁",
                bg="#1e293b",
                fg="#f87171",
                activebackground="#ef4444",
                activeforeground="#ffffff",
                font=("Microsoft JhengHei UI", 8, "bold"),
                relief="solid",
                bd=1,
                padx=9,
                pady=3,
                cursor="hand2",
                command=self.clear_all_company_tabs,
            ).pack(side="left", padx=4, pady=2)

    def remove_company_tab(self, co_id: str):
        self.companies.pop(co_id, None)
        if self.active_tab == co_id:
            self.active_tab = "SUMMARY"
        self.render_chips()
        self.render_summary_table()
        self.switch_tab(self.active_tab)

    def clear_all_company_tabs(self):
        self.companies.clear()
        self.active_tab = "SUMMARY"
        self.render_chips()
        self.render_summary_table()
        self.switch_tab("SUMMARY")
        self.set_status("已清除畫面上所有個股分頁（SQLite 資料庫內資料仍完整保留）。")

    def switch_tab(self, tab_id: str):
        self.active_tab = tab_id
        self.render_tabs()
        self.render_source_info()
        if tab_id == "SUMMARY":
            self.view_company.pack_forget()
            self.view_summary.pack(fill="both", expand=True)
        else:
            self.set_source_co_id(tab_id)
            comp = self.companies.get(tab_id)
            if not comp:
                return
            self.view_summary.pack_forget()
            self.view_company.pack(fill="both", expand=True)

            src_badge = "🗄️ 來源：SQLite 資料庫" if comp.get("source") == "db" else "🌐 來源：MOPS 最新抓取 (已寫入 SQLite)"
            self.lbl_comp_title.configure(text=f"{comp['label']} — 公告自結損益表 ({comp['y1']}年度 vs {comp['y2']}年度)")
            self.lbl_comp_sub.configure(text=f"依照月份（1月至12月）基準對齊 ｜ 單位：新台幣仟元 ｜ {src_badge} ({comp.get('updated_at', '')})")

            self.btn_vc_open_y1.configure(
                text=f"🔗 點擊開啟 MOPS {comp['label']} ({comp['y1']}年度) 原始報表 ↗",
                command=lambda: self.open_mops_company_report(comp["co_id"], comp["y1"]),
            )
            self.btn_vc_open_y2.configure(
                text=f"🔗 點擊開啟 MOPS {comp['label']} ({comp['y2']}年度) 原始報表 ↗",
                command=lambda: self.open_mops_company_report(comp["co_id"], comp["y2"]),
            )
            p_y2 = json.dumps(
                {
                    "apiName": "ajax_t138sb02",
                    "parameters": {"encodeURIComponent": 1, "step": 1, "firstin": 1, "off": 1, "CK2": "1", "BK1": "2", "TYPEK": "sii", "YEAR": str(comp["y2"]), "COMP": str(comp["co_id"])},
                },
                ensure_ascii=False,
            )
            self.lbl_vc_payload.configure(text=f"URL: {API_URL} ｜ Payload: {p_y2}")

            for item in self.tree_comp.get_children():
                self.tree_comp.delete(item)
            for r in comp["rows"]:
                is_ref = r["month"] == self.list_month
                m_lbl = f"{r['month_label']} ⭐" if is_ref else r["month_label"]
                self.tree_comp.insert(
                    "",
                    "end",
                    values=(
                        m_lbl,
                        self.fmt_num(r["y1_month_pretax"]),
                        self.fmt_num(r["y1_cum_pretax"]),
                        self.fmt_num(r["y2_month_pretax"]),
                        self.fmt_num(r["y2_cum_pretax"]),
                        self.fmt_yoy(r["month_yoy"]),
                        self.fmt_yoy(r["cum_yoy"]),
                    ),
                    tags=("ref_month" if is_ref else "normal",),
                )

    # ================================================================ 批次與單獨更新個股
    def update_single_stock(self, co_id: str):
        self.selected.add(co_id)
        self.render_chips()
        self.run_batch_update([co_id])

    def update_all_stocks(self):
        self.select_all(True)
        self.update_selected_stocks()

    def update_selected_stocks(self):
        target_ids = [s["co_id"] for s in self.stocks if s["co_id"] in self.selected]
        if not target_ids:
            messagebox.showinfo("提示", "請先勾選至少一檔金融股，或點擊「⚡ 一鍵更新全部 (ALL)」！")
            return
        self.run_batch_update(target_ids)

    def stop_updating(self):
        self.stop_requested = True
        self.set_status("⏹ 正在停止更新（保留已抓取之個股資料）...")

    def run_batch_update(self, co_ids):
        if self.is_updating:
            return
        self.is_updating = True
        self.stop_requested = False

        self.btn_update_sel.configure(state="disabled")
        self.btn_update_all.configure(state="disabled")
        self.btn_stop.pack(side="left", padx=(0, 6))

        y1 = self.var_y1.get()
        y2 = self.var_y2.get()
        force = self.var_fetch_mode.get() == "force"
        mode_lbl = "強制重抓" if force else "資料庫優先"
        total = len(co_ids)

        def _worker():
            done = 0
            db_hits = 0
            mops_fetches = 0
            try:
                for cid in co_ids:
                    if self.stop_requested:
                        break
                    stock_obj = next((s for s in self.stocks if s["co_id"] == cid), None)
                    label = stock_obj["label"] if stock_obj else cid
                    self.after(0, lambda l=label, d=done, c=cid: (
                        self.set_status(f"[{mode_lbl}] 正在處理 {l} 的 {y1}/{y2} 年度公告自結損益表...", d, total),
                        self.render_summary_table(loading_co_id=c),
                    ))
                    comp = build_company_comparison(cid, y1, y2, ref_month=self.list_month, force=force)
                    self.companies[cid] = comp
                    if comp.get("source") == "db":
                        db_hits += 1
                    else:
                        mops_fetches += 1
                    done += 1
                    self.after(0, lambda l=label, d=done, dh=db_hits, mf=mops_fetches: (
                        self.render_chips(),
                        self.render_summary_table(),
                        self.render_tabs(),
                        self.set_status(f"已完成 {l} ({d}/{total}，資料庫讀取: {dh} 檔，MOPS 重抓寫入: {mf} 檔)", d, total),
                    ))

                if self.stop_requested:
                    self.after(0, lambda d=done: self.set_status(f"⏹ 已手動停止。共完成 {d} / {total} 檔個股更新。", d, total))
                else:
                    self.after(0, lambda d=done, dh=db_hits, mf=mops_fetches: self.set_status(
                        f"✅ 更新完成！共 {d} 檔金融股 (🗄️ 資料庫直接讀取: {dh} 檔 ｜ 🌐 MOPS 抓取寫入: {mf} 檔)。", d, total
                    ))
            except Exception as e:
                err_msg = str(e)
                self.after(0, lambda d=done, m=err_msg: self.set_status(f"❌ 更新過程發生錯誤: {m}", d, total))
            finally:
                self.is_updating = False
                self.after(0, self._on_batch_finished)

        threading.Thread(target=_worker, daemon=True).start()

    def _on_batch_finished(self):
        self.btn_update_sel.configure(state="normal")
        self.btn_update_all.configure(state="normal")
        self.btn_stop.pack_forget()
        self.render_summary_table()

    # ================================================================ 匯出 Excel / CSV
    def export_excel(self):
        y1, y2 = self.var_y1.get(), self.var_y2.get()
        default_name = f"MOPS_金融股自結損益_{self.ym}_{y1}vs{y2}.xlsx"
        filepath = filedialog.asksaveasfilename(
            title="匯出 Excel (.xlsx)",
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[("Excel 活頁簿 (*.xlsx)", "*.xlsx")],
        )
        if not filepath:
            return
        try:
            payload = {
                "ym": self.ym,
                "y1": y1,
                "y2": y2,
                "stocks": self.get_sorted_stocks(),
                "companies": self.companies,
            }
            save_excel_file(filepath, payload)
            self.set_status(f"✅ 已成功匯出 Excel 檔案：{filepath}")
            if messagebox.askyesno("匯出成功", f"已成功匯出 Excel：\n{filepath}\n\n是否立即開啟檔案？"):
                os.startfile(filepath)
        except Exception as e:
            messagebox.showerror("匯出 Excel 失敗", str(e))

    def export_csv(self):
        y1, y2 = self.var_y1.get(), self.var_y2.get()
        default_name = f"MOPS_金融股自結清單_{self.ym}_{y1}vs{y2}.csv"
        filepath = filedialog.asksaveasfilename(
            title="匯出清單 CSV",
            initialfile=default_name,
            defaultextension=".csv",
            filetypes=[("CSV 逗號分隔檔 (*.csv)", "*.csv")],
        )
        if not filepath:
            return
        try:
            ym_lbl = f"{self.list_year}/{self.list_month:02d}"
            rows = []
            for s in self.get_sorted_stocks():
                comp = self.companies.get(s["co_id"])
                latest = comp.get("latest_row") if comp else None
                yoy_val = latest.get("cum_yoy") if latest else None
                rows.append({
                    "金融股": s["label"],
                    "報表類型": s["report_type"],
                    f"{ym_lbl} 本月稅前損益(清單)": s.get("list_month_pretax") if s.get("list_month_pretax") is not None else "",
                    f"{ym_lbl} 本年累計稅前損益(清單)": s.get("list_cum_pretax") if s.get("list_cum_pretax") is not None else "",
                    f"{y1}年度 本月稅前損益": latest.get("y1_month_pretax") if (latest and latest.get("y1_month_pretax") is not None) else "",
                    f"{y1}年度 本年累計稅前損益": latest.get("y1_cum_pretax") if (latest and latest.get("y1_cum_pretax") is not None) else "",
                    f"{y2}年度 本月稅前損益": latest.get("y2_month_pretax") if (latest and latest.get("y2_month_pretax") is not None) else "",
                    f"{y2}年度 本年累計稅前損益": latest.get("y2_cum_pretax") if (latest and latest.get("y2_cum_pretax") is not None) else "",
                    "累計稅前 YoY": f"{yoy_val * 100:.2f}%" if yoy_val is not None else "",
                })
            pd.DataFrame(rows).to_csv(filepath, index=False, encoding="utf-8-sig")
            self.set_status(f"✅ 已成功匯出 CSV 檔案：{filepath}")
        except Exception as e:
            messagebox.showerror("匯出 CSV 失敗", str(e))


if __name__ == "__main__":
    init_db()
    app = MopsDesktopApp()
    app.mainloop()
